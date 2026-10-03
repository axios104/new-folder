"""
Pipeline — connects the automation and scraper subtools (nodriver/async version).

Flow per suburb:
    create_browser()                    -> nd.Browser
    browser.get()                       -> nd.Tab
    navigator.search_suburb()           -> loads results page
    navigator.collect_agent_profile_urls()
                                        -> [profile urls] from ALL pages
    for each url:
        page.get(url)
        extractor.extract_agent_record()
                                        -> raw record
        formatter.cleaners.clean_value() -> cleaned per-field

    all records -> write_output (Excel/JSON) with checkpoints
"""

from __future__ import annotations

import asyncio
import json
import logging
import random
from collections.abc import Callable
from pathlib import Path

import pandas as pd

import config

from automation.browser import create_browser
from automation.navigator import (
    collect_agent_profile_urls,
    search_location,
    search_suburb,
)
from scraper.extractor import extract_agent_record
from scraper.html_fields import extract_team_member_links, split_suburb_and_postcode
from scraper.matching import designation_confidence

logger = logging.getLogger("scraper.pipeline")

CheckpointFn = Callable[[list[dict], list[dict]], None]
RecordFn = Callable[[str, dict], None]


async def _human_pause() -> None:
    await asyncio.sleep(
        random.uniform(
            config.MIN_PROFILE_DELAY_S,
            config.MAX_PROFILE_DELAY_S,
        )
    )


def _record_dedupe_key(row: dict) -> str:
    profile = str(row.get("profile_url", "")).rstrip("/").lower()
    parent = str(row.get("primary_agent_url", "")).rstrip("/").lower()
    return f"{profile}::{parent}" if row.get("record_type") == "Team member" else profile


def _clean_record(raw: dict) -> dict:
    try:
        try:
            from formatter.config import FIELD_BY_KEY
            from formatter.cleaners import clean_value
        except ImportError:
            from recruitment_formatter.formatter.config import FIELD_BY_KEY
            from recruitment_formatter.formatter.cleaners import clean_value

        cleaned = {}
        for key, value in raw.items():
            spec = FIELD_BY_KEY.get(key)
            if spec:
                cleaned[key] = clean_value(spec.dtype, value)
            else:
                cleaned[key] = value
        return cleaned
    except ImportError:
        logger.debug("formatter package not available, skipping cleaning")
        return raw


def _visit_cap(max_agents: int | None) -> int | None:
    if max_agents is not None:
        return max_agents if max_agents > 0 else None
    return config.MAX_AGENTS_PER_SUBURB


async def _scrape_suburbs_async(
    suburbs: list[str],
    headless: bool = False,
    *,
    records: list[dict] | None = None,
    report: list[dict] | None = None,
    on_checkpoint: CheckpointFn | None = None,
    on_record: RecordFn | None = None,
    max_agents: int | None = None,
    designation: str = "",
    deep_search: bool = False,
    use_location_search: bool = False,
    search_terms: dict[str, str] | None = None,
) -> tuple[list[dict], list[dict]]:
    all_records: list[dict] = records if records is not None else []
    report_rows: list[dict] = report if report is not None else []
    seen_urls_by_suburb: dict[str, set[str]] = {}
    for row in all_records:
        query = str(row.get("_suburb_query", ""))
        url = _record_dedupe_key(row)
        if query and url:
            seen_urls_by_suburb.setdefault(query, set()).add(url)

    browser = None
    page = None
    interrupted = False

    def checkpoint(reason: str) -> None:
        logger.info("Saving checkpoint... (%s, %d record(s))", reason, len(all_records))
        if on_checkpoint:
            try:
                on_checkpoint(all_records, report_rows)
                logger.info("Checkpoint saved")
            except Exception as exc:  # noqa: BLE001
                logger.error("Checkpoint failed: %s", exc, exc_info=True)

    try:
        logger.info("Starting nodriver browser (headless=%s)...", headless)
        browser = await create_browser(headless=headless)
        logger.info("Nodriver browser started successfully.")
        page = await browser.get(config.FIND_AGENT_URL)
        logger.info("Browser tab initialized.")

        for suburb in suburbs:
            location_text = (search_terms or {}).get(suburb, suburb)
            suburb_name, postcode = split_suburb_and_postcode(location_text)
            if use_location_search and not suburb_name:
                suburb_name = location_text
            entry = {
                "suburb": suburb,
                "suburb_name": suburb_name,
                "postcode": postcode,
                "profiles_found": 0,
                "pages_scraped": 0,
                "records_scraped": 0,
                "profiles_failed": 0,
                "status": "in_progress",
                "error": None,
            }
            seen_urls = seen_urls_by_suburb.setdefault(suburb, set())

            try:
                logger.info("Processing suburb: %s (suburb=%s postcode=%s)", suburb, suburb_name, postcode)
                ok = await (search_location(page, location_text) if use_location_search else search_suburb(page, suburb))
                if not ok:
                    entry["error"] = "search_failed"
                    entry["status"] = "failed"
                    logger.warning("Search failed for suburb: %s", suburb)
                    report_rows.append(entry)
                    checkpoint(f"search_failed:{suburb_name}")
                    continue

                profile_urls = await collect_agent_profile_urls(
                    page,
                    suburb,
                    on_page=lambda page_no: entry.__setitem__("pages_scraped", page_no),
                ) or []
                entry["profiles_found"] = len(profile_urls)
                logger.info("Found %d agent profile(s) for '%s'.", len(profile_urls), suburb)

                remaining_urls = [
                    url for url in profile_urls
                    if url.rstrip("/").lower() not in seen_urls
                ]
                cap = _visit_cap(max_agents)
                capped = bool(cap and len(remaining_urls) > cap)
                if capped:
                    logger.info(
                        "Visiting first %d of %d collected profile(s) for '%s' (max-agents cap)",
                        cap,
                        len(remaining_urls),
                        suburb,
                    )
                    remaining_urls = remaining_urls[:cap]

                for url in remaining_urls:
                    key = url.rstrip("/").lower()
                    try:
                        logger.info("Opening agent profile: %s", url)
                        await page.get(url)
                        await _human_pause()
                        raw = await extract_agent_record(page, url, suburb_hint=suburb)
                        if not raw:
                            logger.warning("No record extracted from: %s", url)
                            entry["profiles_failed"] += 1
                            continue
                        if not raw.get("suburb"):
                            raw["suburb"] = suburb_name
                        if not raw.get("postcode"):
                            raw["postcode"] = postcode
                        confidence = designation_confidence(str(raw.get("job_title", "")), designation)
                        if designation and confidence < 0.85:
                            logger.info(
                                "Skipping %s: designation %r confidence %.0f%% is below 85%%",
                                raw.get("name") or url,
                                raw.get("job_title", ""),
                                confidence * 100,
                            )
                            continue
                        raw["_suburb_query"] = suburb
                        raw["record_type"] = "Primary agent"
                        raw["primary_agent"] = raw.get("name", "")
                        raw["primary_agent_url"] = url
                        raw["designation_confidence"] = confidence
                        cleaned = _clean_record(raw)
                        all_records.append(cleaned)
                        seen_urls.add(key)
                        entry["records_scraped"] += 1
                        if on_record:
                            try:
                                on_record(suburb, cleaned)
                            except Exception as exc:  # noqa: BLE001
                                entry["profiles_failed"] += 1
                                logger.error("Could not persist profile checkpoint %s: %s", url, exc, exc_info=True)
                        if deep_search and cleaned.get("agency_url"):
                            try:
                                await page.get(cleaned["agency_url"])
                                await _human_pause()
                                agency_html = await page.get_content()
                                team_links = extract_team_member_links(agency_html)
                                logger.info("Found %d team member link(s) for %s", len(team_links), cleaned.get("name", url))
                                for member in team_links:
                                    if member.profile_url.rstrip("/").lower() == url.rstrip("/").lower():
                                        continue
                                    team_row_key = f"{member.profile_url.rstrip('/').lower()}::{url.rstrip('/').lower()}"
                                    if team_row_key in seen_urls:
                                        continue
                                    try:
                                        await page.get(member.profile_url)
                                        await _human_pause()
                                        team_raw = await extract_agent_record(page, member.profile_url, suburb_hint=suburb)
                                        if not team_raw:
                                            entry["profiles_failed"] += 1
                                            continue
                                        team_raw["name"] = team_raw.get("name") or member.name
                                        team_raw["_suburb_query"] = suburb
                                        team_raw["record_type"] = "Team member"
                                        team_raw["primary_agent"] = cleaned.get("name", "")
                                        team_raw["primary_agent_url"] = url
                                        team_raw["designation_confidence"] = ""
                                        team_cleaned = _clean_record(team_raw)
                                        all_records.append(team_cleaned)
                                        seen_urls.add(team_row_key)
                                        entry["records_scraped"] += 1
                                        if on_record:
                                            on_record(suburb, team_cleaned)
                                    except Exception as team_exc:  # noqa: BLE001
                                        entry["profiles_failed"] += 1
                                        logger.warning("Failed team member %s: %s", member.profile_url, team_exc)
                            except Exception as agency_exc:  # noqa: BLE001
                                entry["profiles_failed"] += 1
                                logger.warning("Could not read agency team for %s: %s", url, agency_exc)
                        logger.info(
                            "Scraped: %s (%s / %s) — %s",
                            cleaned.get("name", "?"),
                            cleaned.get("suburb", suburb_name),
                            cleaned.get("postcode", postcode),
                            url,
                        )
                    except KeyboardInterrupt:
                        interrupted = True
                        logger.info("Ctrl+C detected.")
                        logger.info("Stopping scraper gracefully...")
                        break
                    except Exception as exc:  # noqa: BLE001
                        entry["profiles_failed"] += 1
                        logger.error("Failed to scrape %s: %s", url, exc, exc_info=True)
                    await _human_pause()

                if interrupted:
                    entry["status"] = "interrupted"
                    report_rows.append(entry)
                    checkpoint("keyboard_interrupt")
                    break

                entry["status"] = "partial" if entry["profiles_failed"] or capped else "completed"

            except KeyboardInterrupt:
                interrupted = True
                entry["status"] = "interrupted"
                logger.info("Ctrl+C detected.")
                logger.info("Stopping scraper gracefully...")
                report_rows.append(entry)
                checkpoint("keyboard_interrupt")
                break
            except Exception as exc:  # noqa: BLE001
                logger.error("Suburb '%s' failed: %s", suburb, exc, exc_info=True)
                entry["error"] = str(exc)
                entry["status"] = "failed"

            report_rows.append(entry)
            logger.info(
                "Finished suburb '%s': %d profile(s) found, %d record(s) scraped. Total records: %d",
                suburb,
                entry["profiles_found"],
                entry["records_scraped"],
                len(all_records),
            )
            checkpoint(f"suburb:{suburb_name}")
            await _human_pause()

    except KeyboardInterrupt:
        interrupted = True
        logger.info("Ctrl+C detected.")
        logger.info("Stopping scraper gracefully...")
        checkpoint("keyboard_interrupt")
    except Exception as exc:  # noqa: BLE001
        logger.critical("Browser pipeline failed: %s", exc, exc_info=True)
        processed_suburbs = {item["suburb"] for item in report_rows}
        for suburb in suburbs:
            if suburb not in processed_suburbs:
                report_rows.append(
                    {
                        "suburb": suburb,
                        "profiles_found": 0,
                        "records_scraped": 0,
                        "profiles_failed": 0,
                        "status": "failed",
                        "error": f"browser_error: {exc}",
                    }
                )
        checkpoint("browser_error")
    finally:
        if browser is not None:
            try:
                logger.info("Stopping nodriver browser...")
                result = browser.stop()
                if asyncio.iscoroutine(result):
                    await result
                logger.info("Nodriver browser stopped.")
            except Exception as exc:  # noqa: BLE001
                logger.warning("Error while stopping browser: %s", exc)

    return all_records, report_rows


def scrape_suburbs(
    suburbs: list[str],
    headless: bool = False,
    *,
    records: list[dict] | None = None,
    report: list[dict] | None = None,
    on_checkpoint: CheckpointFn | None = None,
    on_record: RecordFn | None = None,
    max_agents: int | None = None,
    designation: str = "",
    deep_search: bool = False,
    use_location_search: bool = False,
    search_terms: dict[str, str] | None = None,
) -> tuple[list[dict], list[dict]]:
    """
    Public synchronous entry point.

    `records` / `report` are filled in place so Ctrl+C cannot drop in-memory data
    even if asyncio.run() is interrupted.
    """
    records = records if records is not None else []
    report = report if report is not None else []
    try:
        return asyncio.run(
            _scrape_suburbs_async(
                suburbs,
                headless=headless,
                records=records,
                report=report,
                on_checkpoint=on_checkpoint,
                on_record=on_record,
                max_agents=max_agents,
                designation=designation,
                deep_search=deep_search,
                use_location_search=use_location_search,
                search_terms=search_terms,
            )
        )
    except KeyboardInterrupt:
        logger.info("Ctrl+C detected.")
        logger.info("Stopping scraper gracefully...")
        if on_checkpoint:
            try:
                on_checkpoint(records, report)
            except Exception as exc:  # noqa: BLE001
                logger.error("Final interrupt checkpoint failed: %s", exc, exc_info=True)
        return records, report


def _schema_columns() -> list[tuple[str, str]]:
    try:
        try:
            from formatter.config import SCHEMA
        except ImportError:
            from recruitment_formatter.formatter.config import SCHEMA
        return [(spec.key, spec.output_name) for spec in SCHEMA]
    except ImportError:
        return [
            ("name", "Name"),
            ("job_title", "job_title"),
            ("years_experience", "Years_experience"),
            ("agency_name", "Agency Name"),
            ("agent_email", "Email ID of the Agent"),
            ("suburb", "suburbs"),
            ("postcode", "Post code"),
            ("agency_address", "agency_address"),
            ("phone", "phone"),
            ("rating", "rating"),
            ("reviews", "Reviews"),
            ("properties_sold", "properties_sold"),
            ("median_sold_price", "median_sold_price"),
            ("median_days_advertised", "median_days_advertised"),
            ("profile_url", "profile_url"),
            ("agency_url", "agency_url"),
        ]


def _records_to_dataframe(records: list[dict]) -> pd.DataFrame:
    columns = _schema_columns()
    relationship_fields = (
        ("record_type", "Record type"),
        ("primary_agent", "Primary agent"),
        ("primary_agent_url", "Primary agent profile"),
        ("designation_confidence", "Designation confidence"),
    )
    rows = []
    for rec in records:
        row = {output: rec.get(key, "") for key, output in columns}
        row.update({output: rec.get(key, "") for key, output in relationship_fields})
        rows.append(row)
    return pd.DataFrame(rows, columns=[output for _, output in columns] + [output for _, output in relationship_fields])


def write_output(
    records: list[dict],
    destination: Path,
    mode: str,
    merge_filename: str = "scraped_data",
) -> Path:
    """
    Writes scraped records to the caller-supplied destination directory.

    Excel mode always writes a real .xlsx file (pandas/openpyxl). It does not
    silently switch to JSON. Missing fields become empty cells.
    """
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    stem = merge_filename or "scraped_data"

    if mode == "excel":
        out_path = destination / f"{stem}.xlsx"
        # Keep the final suffix as .xlsx so both pandas and the shared writer
        # select the correct engine, then atomically replace the prior file.
        tmp_path = destination / f".{stem}.tmp.xlsx"
        df = _records_to_dataframe(records)
        has_relationship_data = any(record.get("record_type") for record in records)
        formatter_written = False
        try:
            try:
                from formatter.writers import write_excel
            except ImportError:
                from recruitment_formatter.formatter.writers import write_excel
            if has_relationship_data:
                with pd.ExcelWriter(tmp_path, engine="openpyxl") as writer:
                    df.to_excel(writer, index=False, sheet_name="Agents")
            else:
                write_excel(records, tmp_path)
            formatter_written = True
        except Exception as exc:  # noqa: BLE001
            logger.warning("Formatter Excel writer unavailable (%s); using pandas/openpyxl fallback", exc)
            with pd.ExcelWriter(tmp_path, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="Agents")
        if has_relationship_data and formatter_written:
            from openpyxl import load_workbook

            workbook = load_workbook(tmp_path)
            worksheet = workbook["Agents"]
            extra_columns = (
                ("record_type", "Record type"),
                ("primary_agent", "Primary agent"),
                ("primary_agent_url", "Primary agent profile"),
                ("designation_confidence", "Designation confidence"),
            )
            first_extra = worksheet.max_column + 1
            for offset, (_, heading) in enumerate(extra_columns):
                column = first_extra + offset
                worksheet.cell(row=1, column=column, value=heading)
                for row_number, record in enumerate(records, start=2):
                    worksheet.cell(row=row_number, column=column, value=record.get(extra_columns[offset][0], ""))
                letter = worksheet.cell(row=1, column=column).column_letter
                max_width = max([len(heading)] + [len(str(record.get(extra_columns[offset][0], ""))) for record in records])
                worksheet.column_dimensions[letter].width = min(max_width + 2, 60)
            workbook.save(tmp_path)
        tmp_path.replace(out_path)
        logger.info("Excel saved to: %s", out_path.resolve())
        return out_path

    out_path = destination / f"{stem}.json"
    try:
        try:
            from formatter.writers import write_json
        except ImportError:
            from recruitment_formatter.formatter.writers import write_json
        write_json(records, out_path)
    except Exception as exc:  # noqa: BLE001
        logger.warning("Formatter JSON writer unavailable (%s); writing JSON directly", exc)
        out_path.write_text(
            json.dumps(records, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
    logger.info("JSON saved to: %s", out_path.resolve())
    return out_path
